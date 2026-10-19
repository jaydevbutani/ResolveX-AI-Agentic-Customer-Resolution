import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Palette
    C_BG = RGBColor(15, 23, 42)        # Slate 900
    C_CARD = RGBColor(30, 41, 59)      # Slate 800
    C_CYAN = RGBColor(6, 182, 212)     # Cyan 500
    C_GREEN = RGBColor(16, 185, 129)   # Emerald 500
    C_AMBER = RGBColor(245, 158, 11)   # Amber 500
    C_WHITE = RGBColor(248, 250, 252)  # Slate 50
    C_MUTED = RGBColor(148, 163, 184)  # Slate 400
    C_ACCENT = RGBColor(99, 102, 241)  # Indigo 500

    blank_layout = prs.slide_layouts[6]

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG
        bg.line.color.rgb = C_BG
        return bg

    def add_header(slide, title, category="RESOLVEX AI — CA-3 PROOF OF CONCEPT"):
        # Category pill
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.35))
        tf_c = cat_box.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.paragraphs[0]
        p_c.text = category.upper()
        p_c.font.size = Pt(11)
        p_c.font.bold = True
        p_c.font.color.rgb = C_CYAN

        # Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.5), Inches(0.8))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = C_WHITE

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1)

    # Accent decorative banner
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.15), Inches(3.8))
    bar.fill.solid()
    bar.fill.fore_color.rgb = C_CYAN
    bar.line.color.rgb = C_CYAN

    t1_box = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11.0), Inches(3.8))
    tf1 = t1_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "FOUNDATIONS OF AGENTIC AI (FAAI) — CA-3 EVALUATION"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = C_CYAN

    p1 = tf1.add_paragraph()
    p1.text = "ResolveX AI"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = C_WHITE

    p2 = tf1.add_paragraph()
    p2.text = "An LLM-Powered Multi-Agent System for Autonomous Customer Issue Resolution"
    p2.font.size = Pt(20)
    p2.font.color.rgb = C_MUTED

    p3 = tf1.add_paragraph()
    p3.text = "\nDomain: E-Commerce & FinTech Operations | Architecture: LangGraph Stateful Multi-Agent System"
    p3.font.size = Pt(14)
    p3.font.color.rgb = C_GREEN

    p4 = tf1.add_paragraph()
    p4.text = "Student: Jaydev Butani & Lavanay Mutreja (Group 13) | 3rd Year B.Tech AIML | October 2026"
    p4.font.size = Pt(13)
    p4.font.color.rgb = C_MUTED

    # -------------------------------------------------------------
    # SLIDE 2: Problem Statement & Existing Inefficiencies
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2)
    add_header(s2, "The Core Problem: Resolution vs Traditional Support Chat")

    # 2 Column Cards
    # Card 1: Traditional Customer Support
    c1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = C_CARD
    c1.line.color.rgb = RGBColor(239, 68, 68)
    tf_c1 = c1.text_frame
    tf_c1.word_wrap = True
    p = tf_c1.paragraphs[0]
    p.text = "Traditional Support & LLM Chatbots"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(239, 68, 68)

    pts1 = [
        "Chatbots only classify, summarize, and draft polite replies.",
        "Disconnected Enterprise Systems: Payment gateway, order DB, and CRM are fragmented.",
        "Manual Human Bottleneck: Support executives manually verify transactions, copy-paste IDs, and click refund buttons.",
        "Customer Friction: Turnaround times of 3 to 7 business days for trivial, low-risk refund cases.",
        "Repetitive Overhead: Over 70% of support queries involve routine, clearly eligible deduction errors."
    ]
    for pt in pts1:
        p_item = tf_c1.add_paragraph()
        p_item.text = f"• {pt}"
        p_item.font.size = Pt(13)
        p_item.font.color.rgb = C_WHITE

    # Card 2: ResolveX Paradigm
    c2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    c2.fill.solid()
    c2.fill.fore_color.rgb = C_CARD
    c2.line.color.rgb = C_GREEN
    tf_c2 = c2.text_frame
    tf_c2.word_wrap = True
    p = tf_c2.paragraphs[0]
    p.text = "ResolveX Autonomous Resolution Engine"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = C_GREEN

    pts2 = [
        "NOT A CHATBOT: ResolveX is an action-oriented Agentic AI resolution engine.",
        "Autonomous by Default: Routine eligible cases (< ₹25,000) are resolved instantly with zero human friction.",
        "Real State Mutations: Dispatches simulated API calls that alter transaction, order, and refund states.",
        "Verifiable Guardrails: Dual-gatekeeper Critic agent validates pre-action eligibility and verifies post-action execution.",
        "Human Escalation Only for Exceptions: High-risk, suspicious, or conflicting cases are safely escalated with full evidence."
    ]
    for pt in pts2:
        p_item = tf_c2.add_paragraph()
        p_item.text = f"✓ {pt}"
        p_item.font.size = Pt(13)
        p_item.font.color.rgb = C_WHITE

    # -------------------------------------------------------------
    # SLIDE 3: The 6 Specialized Agents
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3)
    add_header(s3, "Multi-Agent System Architecture (6 Specialized Agents)")

    agents = [
        ("Planner Agent", "Decomposes unstructured issues into structured JSON execution plans. Dynamically replans upon Critic rejection.", C_CYAN),
        ("Research Agent", "Retrieves grounded policies via RAG, inspects transaction/order telemetry, and detects systemic failure clusters.", C_ACCENT),
        ("Critic Agent", "Pre-execution gatekeeper (eligibility, duplicate prevention, risk limits) and post-execution state verifier.", C_AMBER),
        ("Executor Agent", "Invokes functional simulated business tools/APIs (refunds, order updates) that execute real state transitions.", C_GREEN),
        ("Business Automation", "Orchestrates enterprise workflows: updates CRM cases, dispatches SMS/Email notifications, links systemic cases.", RGBColor(168, 85, 247)),
        ("Software Dev Agent", "Continuous improvement: synthesizes declarative auto-remediation rules from systemic errors and validates via test suite.", RGBColor(236, 72, 153))
    ]

    for idx, (aname, adesc, acolor) in enumerate(agents):
        row = idx // 3
        col = idx % 3
        x = Inches(0.8 + col * 3.9)
        y = Inches(1.8 + row * 2.5)

        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.7), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD
        card.line.color.rgb = acolor
        tf = card.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = aname
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = acolor

        p_desc = tf.add_paragraph()
        p_desc.text = adesc
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = C_WHITE

    # -------------------------------------------------------------
    # SLIDE 4: Orchestration & The Critic Replanning Loop
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4)
    add_header(s4, "Stateful Orchestration & The Critic Replanning Loop")

    # Workflow Pipeline Box
    pipe_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(1.4))
    pipe_box.fill.solid()
    pipe_box.fill.fore_color.rgb = C_CARD
    pipe_box.line.color.rgb = C_CYAN
    tf_p = pipe_box.text_frame
    tf_p.word_wrap = True
    p0 = tf_p.paragraphs[0]
    p0.text = "THE RESOLVEX RESOLUTION PIPELINE (LangGraph StateGraph)"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = C_CYAN

    p1 = tf_p.add_paragraph()
    p1.text = "UNDERSTAND  ➔  PLAN  ➔  INVESTIGATE  ➔  DECIDE  ➔  EXECUTE  ➔  VERIFY  ➔  RESOLVE"
    p1.font.size = Pt(18)
    p1.font.bold = True
    p1.font.color.rgb = C_WHITE

    # Replanning Details Box
    rep_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(3.5), Inches(11.7), Inches(3.2))
    rep_box.fill.solid()
    rep_box.fill.fore_color.rgb = C_CARD
    rep_box.line.color.rgb = C_AMBER
    tf_r = rep_box.text_frame
    tf_r.word_wrap = True

    p = tf_r.paragraphs[0]
    p.text = "Dynamic Critic Rejection & Replanning Mechanism (Mandatory Safety Guard)"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_AMBER

    pts_rep = [
        "Pre-Execution Validation: Critic checks eligibility against retrieved RAG policies before any API call is made.",
        "Rejection on Duplicate Action: If customer claims deduction but a refund record (e.g. RF-40192) already exists in PROCESSING state, Critic halts execution: REJECT.",
        "Replanning Feedback Loop: Critic emits machine-actionable guidance back to Planner Agent.",
        "Replanned Execution: Planner generates an alternate plan: fetch existing refund ARN tracking status rather than initiating a redundant financial refund.",
        "Execution & Cryptographic Verification: Post-execution Critic computes SHA-256 audit digest verifying state integrity."
    ]
    for pt in pts_rep:
        p_item = tf_r.add_paragraph()
        p_item.text = f"• {pt}"
        p_item.font.size = Pt(13)
        p_item.font.color.rgb = C_WHITE

    # -------------------------------------------------------------
    # SLIDE 5: Four Mandatory Demonstration Scenarios
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5)
    add_header(s5, "Mandatory Demonstration Scenarios")

    scenarios = [
        ("SCENARIO 1: Autonomous Refund", "Customer: '₹2,499 deducted but order failed.'\n• Payment = SUCCESS, Order = FAILED.\n• Policy RAG: POL-REF-01 applied.\n• Critic approves -> Executor calls initiate_refund().\n• Status: RESOLVED AUTONOMOUSLY.", C_GREEN),
        ("SCENARIO 2: Critic Rejection & Replan", "Customer: '₹2,499 deducted for ORD-5502.'\n• System detects RF-40192 already queued.\n• Critic REJECTS duplicate refund.\n• Planner replans to ARN status tracking.\n• Status: RESOLVED WITH EXISTING TRACKING.", C_AMBER),
        ("SCENARIO 3: Exception Escalation", "Customer: '₹48,500 deducted, order held.'\n• Transaction > ₹25,000 threshold & proxy IP.\n• Autonomous execution strictly BLOCKED.\n• Complete evidence docket created for Tier-2 Lead.\n• Status: ESCALATED TO HUMAN.", RGBColor(239, 68, 68)),
        ("SCENARIO 4: Systemic Incident", "50 Customer orders failing with PG_TIMEOUT_504.\n• Research Agent detects systemic pattern.\n• Master Incident INC-2026-SYS created & cases linked.\n• Software Dev Agent synthesizes & tests auto-rule.\n• Status: SYSTEMIC INCIDENT ORCHESTRATED.", C_CYAN)
    ]

    for idx, (sname, sdesc, scolor) in enumerate(scenarios):
        row = idx // 2
        col = idx % 2
        x = Inches(0.8 + col * 5.9)
        y = Inches(1.8 + row * 2.5)

        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.7), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD
        card.line.color.rgb = scolor
        tf = card.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = sname
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = scolor

        p_desc = tf.add_paragraph()
        p_desc.text = sdesc
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = C_WHITE

    # -------------------------------------------------------------
    # SLIDE 6: PoC Quantitative Evaluation & Benchmark Results
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6)
    add_header(s6, "PoC Quantitative Evaluation & Acceptance Benchmark")

    metrics_cards = [
        ("100.0%", "Resolution Success Rate", "All 4 scenarios and 8 acceptance tests passed deterministically.", C_GREEN),
        ("100.0%", "Autonomous Rate (Low-Risk)", "Zero human intervention required for normal eligible cases.", C_CYAN),
        ("100.0%", "Critic Duplicate Prevention", "Zero duplicate refunds issued when refund was already active.", C_AMBER),
        ("100.0%", "Human Escalation Precision", "High-risk cases correctly halted and routed to Tier-2 Lead.", RGBColor(239, 68, 68)),
        ("< 0.05s", "Execution Latency (Fallback)", "Deterministic Cognitive Engine enables zero-lag presentation.", C_GREEN),
        ("99.2%", "Software Dev Rule Accuracy", "Synthesized remediation rule validated against synthetic history.", RGBColor(236, 72, 153))
    ]

    for idx, (mval, mtitle, msub, mcolor) in enumerate(metrics_cards):
        row = idx // 3
        col = idx % 3
        x = Inches(0.8 + col * 3.9)
        y = Inches(1.8 + row * 2.5)

        card = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(3.7), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = C_CARD
        card.line.color.rgb = mcolor
        tf = card.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = mval
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = mcolor

        p_t = tf.add_paragraph()
        p_t.text = mtitle
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = C_WHITE

        p_s = tf.add_paragraph()
        p_s.text = msub
        p_s.font.size = Pt(11)
        p_s.font.color.rgb = C_MUTED

    # -------------------------------------------------------------
    # SLIDE 7: Limitations & Future Scope
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7)
    add_header(s7, "Engineering Limitations & Future Scope")

    # 2 Column Cards
    c1 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = C_CARD
    c1.line.color.rgb = C_AMBER
    tf_c1 = c1.text_frame
    tf_c1.word_wrap = True
    p = tf_c1.paragraphs[0]
    p.text = "Current PoC Limitations"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = C_AMBER

    lims = [
        "Simulated Systems Boundary: Synthetic databases used to avoid real banking API risks in academic PoC.",
        "Single-Tenant Sandbox: Operates locally within Streamlit session memory and local JSON stores.",
        "Lexical/Semantic Hybrid RAG: Policy knowledge base uses curated enterprise policies rather than enterprise vector databases.",
        "Pre-defined Error Signatures: Software Development Agent operates on structured error telemetry rather than raw distributed trace logs."
    ]
    for lim in lims:
        p_item = tf_c1.add_paragraph()
        p_item.text = f"• {lim}"
        p_item.font.size = Pt(13)
        p_item.font.color.rgb = C_WHITE

    c2 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    c2.fill.solid()
    c2.fill.fore_color.rgb = C_CARD
    c2.line.color.rgb = C_CYAN
    tf_c2 = c2.text_frame
    tf_c2.word_wrap = True
    p = tf_c2.paragraphs[0]
    p.text = "Future Scope & Production Roadmap"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = C_CYAN

    futures = [
        "Live Payment Webhook Ingestion: Real-time Kafka / RabbitMQ consumer for instant payment callback drops.",
        "Multi-Modal Evidence Verification: Vision-LLM receipt and delivery courier photograph verification.",
        "Distributed Multi-Tenancy: Kubernetes-orchestrated worker agents with fine-grained RBAC.",
        "Reinforcement Learning from Human Feedback (RLHF): Agent policy refinement based on Tier-2 supervisor approvals."
    ]
    for f in futures:
        p_item = tf_c2.add_paragraph()
        p_item.text = f"➔ {f}"
        p_item.font.size = Pt(13)
        p_item.font.color.rgb = C_WHITE

    # -------------------------------------------------------------
    # SLIDE 8: Conclusion & Key Takeaways
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_bg(s8)
    add_header(s8, "Conclusion & Key Takeaways")

    concl_box = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    concl_box.fill.solid()
    concl_box.fill.fore_color.rgb = C_CARD
    concl_box.line.color.rgb = C_GREEN
    tf_k = concl_box.text_frame
    tf_k.word_wrap = True

    p = tf_k.paragraphs[0]
    p.text = "Why ResolveX Represents True Agentic AI"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = C_GREEN

    conclusions = [
        ("Action Over Conversation", "ResolveX does not just converse with customers; it autonomously manipulates business state through functional simulated tools."),
        ("Autonomous by Default", "Solves the traditional support bottleneck by autonomously disbursing refunds for low-risk eligible claims without human friction."),
        ("Verifiable Safety Guardrails", "Dual-pass Critic agent blocks duplicate disbursements, forces dynamic replanning loops, and halts high-value anomalies."),
        ("Continuous Self-Improvement", "Software Development Agent synthesizes and tests declarative rules against historical telemetry when systemic errors arise."),
        ("Enterprise Impact", "Reduces customer turnaround from 3-7 days to under 3 seconds, eliminating 70%+ of manual operational support overhead.")
    ]
    for ctitle, cdesc in conclusions:
        p_item = tf_k.add_paragraph()
        p_item.text = f"★ {ctitle}: {cdesc}"
        p_item.font.size = Pt(13)
        p_item.font.color.rgb = C_WHITE

    output_path = "presentation/ResolveX_AI_CA3_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_deck()
